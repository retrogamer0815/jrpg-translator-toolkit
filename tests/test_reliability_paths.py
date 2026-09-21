"""Offline regressions: synthetic DBs/media, no provider, Anki, or device I/O."""
import ast
import asyncio
import base64
from contextlib import closing, redirect_stdout
from datetime import datetime
import hashlib
import io
import json
import mimetypes
import os
from pathlib import Path
import re
import shutil
import sqlite3
import sys
import tempfile
from types import SimpleNamespace
from typing import List, Tuple
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import anki_candidates as candidates
import study_library as library
import study_library_sections as sections
import model_catalog as catalog


def extract(filename, names, namespace=None):
    tree = ast.parse((ROOT / "scripts" / filename).read_text(encoding="utf-8-sig"))
    nodes = [n for n in tree.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)) and n.name in names]
    if len(nodes) != len(names):
        raise AssertionError("Missing production functions: " + filename)
    ns = dict(namespace or {})
    exec(compile(ast.Module(body=nodes, type_ignores=[]), filename, "exec"), ns)
    return ns


class ReliabilityPaths(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="jrpg-reliability-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def archive(self, namespace=None):
        if namespace is None:
            tree = ast.parse((ROOT / "scripts/explainer.py").read_text(encoding="utf-8-sig"))
            schema = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
                          and any(isinstance(t, ast.Name) and t.id == "STUDY_LIBRARY_SCHEMA" for t in n.targets))
            ns = dict(globals(), STUDY_LIBRARY_SCHEMA=schema + sections.SECTION_SCHEMA,
                      ensure_section_schema=sections.ensure_section_schema,
                      replace_explanation_sections=sections.replace_explanation_sections,
                      extract_strict_speaker_header=sections.extract_strict_speaker_header)
            namespace = extract("explainer.py", {"initialize_study_library", "archive_study_library_entry",
                                "study_source_identity", "file_sha256"}, ns)
        from PIL import Image
        source = self.root / "source.png"
        Image.new("RGB", (8, 8), "blue").save(source)
        with redirect_stdout(io.StringIO()):
            namespace["archive_study_library_entry"](
                study_dir=str(self.root / "library"), jp="「Test」\n日本語", source_paths=[str(source)],
                text="Original Japanese:\n日本語\n\nDetailed analysis:\nTest.",
                game_profile="fixture", provider="fixture", model="fixture", prompt_profile="fixture",
                key_grammar="", chapter="", include_screenshots=True)
        return self.root / "library/study_library.db", namespace

    def test_read_only_paths_are_escaped_and_never_create_databases(self):
        export_reader = extract("study_export.py", {"connect_read_only"}, dict(Path=Path, sqlite3=sqlite3))["connect_read_only"]
        readers = [library.connect_read_only, candidates.connect_read_only, export_reader]
        for name in ["Game#1", "Game%23", "Game space 日本語"]:
            folder = self.root / name
            folder.mkdir()
            database = folder / "study_library.db"
            with closing(sqlite3.connect(database)) as db:
                db.execute("CREATE TABLE sentinel(value TEXT)")
                db.execute("INSERT INTO sentinel VALUES ('correct library')")
                db.commit()
            before = sorted(str(p.relative_to(self.root)) for p in self.root.rglob("*"))
            for reader in readers:
                with closing(reader(database)) as db:
                    self.assertEqual(db.execute("SELECT value FROM sentinel").fetchone()[0], "correct library")
                    with self.assertRaises(sqlite3.OperationalError):
                        db.execute("CREATE TABLE unwanted(x)")
                with self.assertRaises(sqlite3.OperationalError):
                    reader(folder / "missing#db.sqlite")
            self.assertEqual(before, sorted(str(p.relative_to(self.root)) for p in self.root.rglob("*")))

    def test_excel_source_strings_are_literal_but_dates_and_numbers_are_typed(self):
        try:
            import study_export as export
            from openpyxl import Workbook, load_workbook
        except (ImportError, OSError) as exc:
            self.skipTest("Workbook dependencies unavailable in this runtime: " + str(exc))
        workbook = Workbook()
        export.add_sheet(workbook, "Test", ["Profile", "Japanese", "Count", "Date generated"],
                         [["=1+1", "=HYPERLINK(\"https://example.invalid\",\"日本語\")", 3, datetime(2026, 1, 1)]])
        target = self.root / "export.xlsx"
        workbook.save(target)
        with zipfile.ZipFile(target) as archive:
            self.assertNotIn(b"<f>", archive.read("xl/worksheets/sheet2.xml"))
        with closing(load_workbook(target, data_only=False)) as saved:
            self.assertEqual(saved["Test"]["A2"].value, "=1+1")
            self.assertEqual(saved["Test"]["A2"].data_type, "s")
            self.assertEqual(saved["Test"]["B2"].data_type, "s")
            self.assertEqual(saved["Test"]["C2"].data_type, "n")
            self.assertEqual(saved["Test"]["D2"].data_type, "d")

    def recommendation_fixture(self):
        ids = [hashlib.sha256(str(n).encode()).hexdigest() for n in range(37)]
        rows = [["", "", "", candidates.encode_field("日本語 " + str(n)), "", key, "-1", "0", ""]
                for n, key in enumerate(ids)]
        (self.root / "candidate_sentences.tsv").write_text("\n".join("\t".join(r) for r in rows), encoding="utf-8")
        return ids, self.root / "preferences.db"

    @staticmethod
    def response(ids):
        return json.dumps({"items": [dict(id=key, recommended=True, score=4, reason="Fixture") for key in ids]})

    def test_later_batch_failure_checkpoints_and_stale_snapshot_retry_resumes(self):
        ids, prefs = self.recommendation_fixture()
        # A hallucinated ID from the next batch must not get cached prematurely.
        with patch.object(candidates, "call_model", side_effect=[self.response(ids), TimeoutError("batch 2 failed")]):
            self.assertEqual(candidates.generate_recommendations(self.root, prefs, "openai", "fixture"), 2)
        with closing(sqlite3.connect(prefs)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM " + candidates.RECOMMENDATION_TABLE).fetchone()[0], 36)
        self.assertIn("36 completed assessments were saved", (self.root / "candidate_recommendation_error.txt").read_text())
        with patch.object(candidates, "call_model", return_value=self.response(ids[36:])) as model:
            self.assertEqual(candidates.generate_recommendations(self.root, prefs, "openai", "fixture"), 0)
        model.assert_called_once()
        prompt = model.call_args.args[2]
        self.assertIn(ids[36], prompt)
        self.assertNotIn(ids[0], prompt)
        with closing(sqlite3.connect(prefs)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM " + candidates.RECOMMENDATION_TABLE).fetchone()[0], 37)

    def test_cancel_preserves_completed_batches(self):
        ids, prefs = self.recommendation_fixture()
        with patch.object(candidates, "call_model", side_effect=[self.response(ids[:36]), KeyboardInterrupt()]):
            with self.assertRaises(KeyboardInterrupt):
                candidates.generate_recommendations(self.root, prefs, "openai", "fixture")
        with closing(sqlite3.connect(prefs)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM " + candidates.RECOMMENDATION_TABLE).fetchone()[0], 36)

    def test_archive_postcommit_log_failure_keeps_committed_media(self):
        database, ns = self.archive()
        def broken_log(*args, **kwargs): raise BrokenPipeError("closed log")
        with patch.dict(ns, print=broken_log):
            self.archive(ns)
        with closing(sqlite3.connect(database)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM explanations").fetchone()[0], 2)
            for row in db.execute("SELECT relative_path FROM media"):
                self.assertTrue((database.parent / row[0]).is_file())

    def test_archive_precommit_failure_rolls_back(self):
        database, ns = self.archive()
        with patch.object(shutil, "copy2", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError): self.archive(ns)
        with closing(sqlite3.connect(database)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM explanations").fetchone()[0], 1)
            self.assertEqual(db.execute("PRAGMA foreign_key_check").fetchall(), [])
        self.assertEqual(len(list((database.parent / "Media").glob("*.png"))), 1)

    def test_delete_response_failure_keeps_committed_deletion_and_backup(self):
        database, _ = self.archive()
        with patch.object(library, "write_rows", side_effect=PermissionError("status locked")):
            with self.assertRaisesRegex(RuntimeError, "deletion was committed"):
                library.remove_version(database, self.root / "delete", 1, 1)
        with closing(sqlite3.connect(database)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM explanations").fetchone()[0], 0)
        self.assertFalse((database.parent / "Media/00000001_01.png").exists())
        self.assertEqual(len(list((database.parent / "Trash").rglob("*.png"))), 1)
        self.assertEqual(len(list((database.parent / "Backups").glob("*.db"))), 1)

    def test_delete_precommit_failure_restores_database_and_media(self):
        database, _ = self.archive()
        with patch.object(shutil, "move", side_effect=PermissionError("media locked")):
            with self.assertRaises(PermissionError): library.remove_version(database, self.root / "delete", 1, 1)
        with closing(sqlite3.connect(database)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM explanations").fetchone()[0], 1)
        self.assertTrue((database.parent / "Media/00000001_01.png").is_file())

    def test_malformed_cache_can_be_refreshed(self):
        invalid = [[], None, "string", {}, {"schema_version": catalog.SCHEMA_VERSION, "fetched_at": "2026-01-01", "models": [None]},
                   {"schema_version": catalog.SCHEMA_VERSION, "fetched_at": "2026-01-01", "models": [{"id": "test", "supported_actions": None}]}]
        for value in invalid:
            (self.root / "openai_models.json").write_text(json.dumps(value), encoding="utf-8")
            with patch.object(catalog, "_fetch_openai", return_value=[]) as fetch:
                self.assertTrue(catalog._catalog("openai", "all", self.root, True, 24)["ok"])
                fetch.assert_called_once()

    def test_audio_sessions_supervise_failure_closure_and_cancellation(self):
        async def scenario(provider, outcome):
            sender_started = asyncio.Event()
            receiver_started = asyncio.Event()
            cleaned = set()
            class Socket:
                async def __aenter__(self): return self
                async def __aexit__(self, *args): cleaned.add("socket")
                async def send(self, value): pass
                async def recv(self): return '{"setupComplete":{}}'
                def __aiter__(self): return self
                async def __anext__(self):
                    receiver_started.set()
                    try:
                        if outcome == "receiver": raise ConnectionError("receive failed")
                        if outcome == "closed": raise StopAsyncIteration
                        await asyncio.Event().wait()
                    finally: cleaned.add("receiver")
            async def sender(*args):
                sender_started.set()
                try:
                    await receiver_started.wait()
                    if outcome == "sender": raise OSError("device disconnected")
                    if outcome == "ended": return
                    await asyncio.Event().wait()
                finally: cleaned.add("sender")
            ns = dict(asyncio=asyncio, json=json, base64=base64,
                      websockets=SimpleNamespace(connect=lambda *a, **kw: Socket()),
                      OPENAI_API_KEY="fixture", GOOGLE_API_KEY="fixture", TRANSLATE_MODEL="fixture",
                      GEMINI_AUDIO_MODEL="fixture", TARGET_LANGUAGE_CODE="en", DEBUG=False,
                      _mark_audio_connection_ready=lambda: None, TranscriptBuffer=lambda: None,
                      audio_sender=sender, log=lambda *a: None)
            ns = extract("live_audio_translator.py", {"run_openai", "run_gemini", "_decode_ws_text", "supervise_audio_session"}, ns)
            task = asyncio.create_task(ns[provider](object()))
            if outcome == "cancel":
                await asyncio.wait_for(sender_started.wait(), 1)
                await asyncio.wait_for(receiver_started.wait(), 1)
                task.cancel()
                with self.assertRaises(asyncio.CancelledError): await task
            elif outcome == "closed": await asyncio.wait_for(task, 1)
            else:
                with self.assertRaises((OSError, RuntimeError)):
                    await asyncio.wait_for(task, 1)
            self.assertEqual(cleaned, {"sender", "receiver", "socket"})
        for provider in ("run_openai", "run_gemini"):
            for outcome in ("sender", "receiver", "closed", "cancel", "ended"):
                with self.subTest(provider=provider, outcome=outcome): asyncio.run(scenario(provider, outcome))

    def test_capture_end_and_device_errors_are_actionable_not_network_retries(self):
        async def scenario(mode):
            def blocks(speaker):
                if mode == "device": raise OSError("device unplugged")
                if mode == "ended": return
                yield b"samples"
            class Socket:
                async def send(self, text): raise ConnectionError("network disconnected")
            ns = extract("live_audio_translator.py", {"audio_sender", "AudioCaptureError", "_is_retryable_connection_error"},
                         dict(asyncio=asyncio, json=json, capture_blocks=blocks))
            with self.assertRaises(Exception) as caught:
                await asyncio.wait_for(ns["audio_sender"](Socket(), object(), lambda b: {}), 1)
            if mode == "network": self.assertIsInstance(caught.exception, ConnectionError)
            else:
                self.assertIsInstance(caught.exception, ns["AudioCaptureError"])
                self.assertFalse(ns["_is_retryable_connection_error"](caught.exception))
                self.assertIn("Windows output device", str(caught.exception))
        for mode in ("device", "ended", "network"):
            with self.subTest(mode=mode): asyncio.run(scenario(mode))


if __name__ == "__main__":
    unittest.main()
